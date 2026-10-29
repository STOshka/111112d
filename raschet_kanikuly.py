# -*- coding: utf-8 -*-
"""
Расчёт рабочего времени педагогических работников в период каникул.

Совместимо с Python 3.8 (последняя версия Python для Windows 7).
Собирается в один .exe через PyInstaller (см. build_exe.bat).

Формула:
    время работы в день = недельная нагрузка / число рабочих дней в неделю,
                          округлено ВВЕРХ до целой минуты
    время ухода        = время прихода + время работы в день + обеденный перерыв
Все расчёты ведутся в целых минутах через Decimal — без ошибок округления.
"""

import base64
import datetime
import hashlib
import hmac
import json
import os
import re
import stat
import sys
import zlib
from decimal import Decimal, InvalidOperation, ROUND_CEILING, ROUND_HALF_UP

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

APP_TITLE = "Расчёт рабочего времени в период каникул"

DATA_FILE_NAME = "raschet_kanikuly.dat"            # автосохранение (защищённое)
DATA_EXT = ".dat"
CONFIG_FILE_NAME = "raschet_kanikuly.cfg"         # последний открытый файл и недавние
MAX_RECENT = 6
DATA_HEADER = "RASCHET-KANIKULY-DATA-2"
# Ключ подписи файла данных. Любое ручное изменение файла обнаруживается.
_SIGN_KEY = b"RK-2026-\x8f\x13\xa7kanikuly\x5c\xe2\x91-podpis-\x04\x7d"

MAX_LOAD = Decimal("60")          # максимальная недельная нагрузка, ч
LATE_WARNING_MIN = 20 * 60        # уход позже 20:00 — подсветить жёлтым
DAY_END_MIN = 24 * 60             # уход в 24:00 и позже — ошибка
DAYS_CHOICES = [str(d) for d in range(1, 8)]
ROUNDING_TEXT = "вверх до целой минуты"
DEFAULT_NOTE = "в соответствии с утверждённым графиком"
MAX_NOTE_LEN = 200


# ----------------------------------------------------------------------------
# Разбор и проверка ввода
# ----------------------------------------------------------------------------

class InputError(ValueError):
    """Ошибка ввода с понятным пользователю текстом."""


def normalize_name(text):
    return " ".join(str(text).split())


def parse_name(text):
    name = normalize_name(text)
    if not name:
        raise InputError("Введите Ф.И.О. работника.")
    if re.search(r"\d", name):
        raise InputError("Ф.И.О. не должно содержать цифр.")
    if len(name) > 120:
        raise InputError("Слишком длинное Ф.И.О.")
    return name


def parse_load(text):
    """Недельная нагрузка: число > 0, не больше MAX_LOAD, до 2 знаков после запятой."""
    s = str(text).strip().replace(" ", "").replace(",", ".")
    if not s:
        raise InputError("Введите недельную нагрузку (часы).")
    if not re.fullmatch(r"\d+(\.\d+)?", s):
        raise InputError("Нагрузка должна быть числом, например 28,5")
    try:
        value = Decimal(s)
    except InvalidOperation:
        raise InputError("Нагрузка должна быть числом, например 28,5")
    if value <= 0:
        raise InputError("Нагрузка должна быть больше нуля.")
    if value > MAX_LOAD:
        raise InputError("Нагрузка не может быть больше {} ч в неделю.".format(MAX_LOAD))
    if value != value.quantize(Decimal("0.01")):
        raise InputError("Нагрузка — не больше двух знаков после запятой.")
    return value.quantize(Decimal("0.01"))


def parse_time(text):
    """Время 'ч:мм' (также 'ч.мм', 'ч,мм', 'ч') -> минуты от полуночи."""
    s = str(text).strip()
    if not s:
        raise InputError("Введите время в формате ч:мм, например 8:30")
    m = re.fullmatch(r"(\d{1,2})(?:\s*[:.,]\s*(\d{2}))?", s)
    if not m:
        raise InputError("Время вводится в формате ч:мм, например 8:30")
    hours = int(m.group(1))
    minutes = int(m.group(2) or 0)
    if hours > 23:
        raise InputError("Часы должны быть от 0 до 23.")
    if minutes > 59:
        raise InputError("Минуты должны быть от 00 до 59.")
    return hours * 60 + minutes


def parse_date(text, what):
    s = str(text).strip()
    try:
        return datetime.datetime.strptime(s, "%d.%m.%Y").date()
    except ValueError:
        raise InputError("Дата «{}» должна быть в формате ДД.ММ.ГГГГ, например 01.11.2026".format(what))


def parse_int_range(text, what, lo, hi):
    s = str(text).strip()
    if not re.fullmatch(r"\d+", s):
        raise InputError("{}: введите целое число.".format(what))
    v = int(s)
    if not lo <= v <= hi:
        raise InputError("{}: допустимо от {} до {}.".format(what, lo, hi))
    return v


def parse_note(text):
    note = " ".join(str(text).split())
    if not note:
        raise InputError("Введите текст, который будет вместо расчёта.")
    if len(note) > MAX_NOTE_LEN:
        raise InputError("Текст слишком длинный (не больше {} символов).".format(MAX_NOTE_LEN))
    return note


def parse_days(text):
    return parse_int_range(text, "Рабочих дней в неделю", 1, 7)


# ----------------------------------------------------------------------------
# Форматирование
# ----------------------------------------------------------------------------

def fmt_time(minutes):
    return "{}:{:02d}".format(minutes // 60, minutes % 60)


fmt_duration = fmt_time


def fmt_decimal(value, places=2):
    q = Decimal(1).scaleb(-places)
    return str(Decimal(value).quantize(q, rounding=ROUND_HALF_UP)).replace(".", ",")


def fmt_load(value):
    s = "{:.2f}".format(Decimal(value))
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s.replace(".", ",")


def days_word(n):
    if n % 10 == 1 and n % 100 != 11:
        return "день"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return "дня"
    return "дней"


# ----------------------------------------------------------------------------
# Модель данных и расчёт
# ----------------------------------------------------------------------------

class Settings(object):
    def __init__(self, date_from=None, date_to=None, days=5, lunch=20,
                 default_arrival=8 * 60 + 30):
        today = datetime.date.today()
        self.date_from = date_from or today
        self.date_to = date_to or today
        self.days = days                    # общее число рабочих дней в неделю
        self.lunch = lunch
        self.default_arrival = default_arrival

    def to_dict(self):
        return {
            "date_from": self.date_from.strftime("%d.%m.%Y"),
            "date_to": self.date_to.strftime("%d.%m.%Y"),
            "days": self.days,
            "lunch": self.lunch,
            "default_arrival": fmt_time(self.default_arrival),
        }

    @classmethod
    def from_dict(cls, d):
        s = cls()
        try:
            s.date_from = parse_date(d.get("date_from", ""), "с")
            s.date_to = parse_date(d.get("date_to", ""), "по")
        except InputError:
            pass
        for attr, parser in (("days", parse_days),
                             ("lunch", lambda v: parse_int_range(v, "Обед", 0, 180)),
                             ("default_arrival", parse_time)):
            try:
                setattr(s, attr, parser(d.get(attr, getattr(s, attr) if attr != "default_arrival" else "8:30")))
            except InputError:
                pass
        return s


class Employee(object):
    def __init__(self, name, load, arrival, days=None, note=None):
        self.name = name          # str
        self.load = load          # Decimal, часов в неделю (может быть None, если есть note)
        self.arrival = arrival    # int, минуты от полуночи (может быть None, если есть note)
        self.days = days          # None = как в общих параметрах, иначе 1..7
        self.note = note          # текст на всю строку вместо расчёта, например
                                  # «в соответствии с утверждённым графиком»
        if note is None and (load is None or arrival is None):
            raise InputError("Нужны нагрузка и время прихода (или текст вместо расчёта).")

    def days_for(self, settings):
        return self.days if self.days else settings.days

    def to_dict(self):
        return {"name": self.name,
                "load": "{:.2f}".format(self.load) if self.load is not None else "",
                "arrival": fmt_time(self.arrival) if self.arrival is not None else "",
                "days": self.days, "note": self.note}

    @classmethod
    def from_dict(cls, d):
        days = d.get("days")
        note = d.get("note")
        note = parse_note(note) if note else None
        load = d.get("load") or ""
        arr = d.get("arrival") or ""
        return cls(parse_name(d["name"]),
                   parse_load(load) if (load or not note) else None,
                   parse_time(arr) if (arr or not note) else None,
                   parse_days(days) if days not in (None, "", 0) else None,
                   note)


class Result(object):
    def __init__(self, days, work_hours, work_min, total_min, departure):
        self.days = days
        self.work_hours = work_hours    # Decimal — точное значение в часах
        self.work_min = work_min        # int — работа в день, минут (округлено вверх)
        self.total_min = total_min      # int — с обедом
        self.departure = departure      # int — минуты от полуночи

    @property
    def is_error(self):
        return self.departure >= DAY_END_MIN

    @property
    def is_late(self):
        return LATE_WARNING_MIN < self.departure < DAY_END_MIN


def calculate(load, arrival, settings, days=None):
    days = days or settings.days
    load = Decimal(load)
    work_hours = load / Decimal(days)
    exact_minutes = load * 60 / Decimal(days)
    work_min = int(exact_minutes.quantize(Decimal(1), rounding=ROUND_CEILING))
    total_min = work_min + settings.lunch
    return Result(days, work_hours, work_min, total_min, arrival + total_min)


def calc_employee(e, settings):
    """Результат расчёта или None, если вместо расчёта текст."""
    if e.note:
        return None
    return calculate(e.load, e.arrival, settings, e.days)


# ----------------------------------------------------------------------------
# Защищённый файл данных
# ----------------------------------------------------------------------------
# Формат: строка-заголовок, затем JSON, сжатый zlib и закодированный base64,
# затем подпись HMAC-SHA256. Открыть в Блокноте и поправить цифры нельзя:
# текст нечитаем, а любое изменение ломает подпись, и программа такой файл
# не примет.

class DataFileError(Exception):
    pass


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


def _dir_writable(path):
    try:
        test = os.path.join(path, ".write_test.tmp")
        with open(test, "w") as f:
            f.write("ok")
        os.remove(test)
        return True
    except OSError:
        return False


def data_dir():
    d = app_dir()
    if not _dir_writable(d):
        base = os.environ.get("APPDATA") or os.path.expanduser("~")
        d = os.path.join(base, "RaschetKanikuly")
        try:
            os.makedirs(d)
        except OSError:
            pass
    return d


def _sign(payload):
    return hmac.new(_SIGN_KEY, payload, hashlib.sha256).hexdigest()


def _make_writable(path):
    if os.path.exists(path):
        try:
            os.chmod(path, stat.S_IREAD | stat.S_IWRITE)
        except OSError:
            pass


def save_data(path, settings, employees):
    data = {
        "app": "RaschetKanikuly",
        "saved": datetime.datetime.now().strftime("%d.%m.%Y %H:%M"),
        "settings": settings.to_dict(),
        "employees": [e.to_dict() for e in employees],
    }
    raw = json.dumps(data, ensure_ascii=False, sort_keys=True).encode("utf-8")
    payload = base64.b64encode(zlib.compress(raw, 9))
    text = "{}\n{}\n{}\n".format(DATA_HEADER, payload.decode("ascii"), _sign(payload))
    tmp = path + ".tmp"
    _make_writable(tmp)
    with open(tmp, "w", encoding="ascii", newline="\n") as f:
        f.write(text)
    _make_writable(path)
    if os.path.exists(path):
        os.remove(path)
    os.rename(tmp, path)
    try:
        os.chmod(path, stat.S_IREAD)  # атрибут «Только чтение»
    except OSError:
        pass


def _parse_data_dict(data):
    settings = Settings.from_dict(data.get("settings", {}))
    employees, bad = [], 0
    for d in data.get("employees", []):
        try:
            employees.append(Employee.from_dict(d))
        except (InputError, KeyError, TypeError, AttributeError):
            bad += 1
    return settings, employees, bad


def load_data(path):
    """Читает защищённый файл. DataFileError — если файл изменён или повреждён."""
    try:
        with open(path, "r", encoding="ascii") as f:
            lines = f.read().split()
    except (OSError, UnicodeDecodeError):
        raise DataFileError("Файл не читается или это не файл данных программы.")
    if len(lines) != 3 or lines[0] != DATA_HEADER:
        raise DataFileError("Это не файл данных программы (или он был изменён вручную).")
    payload, sig = lines[1].encode("ascii"), lines[2]
    if not hmac.compare_digest(_sign(payload), sig):
        raise DataFileError("Файл был изменён вручную или повреждён — подпись не совпадает.")
    try:
        data = json.loads(zlib.decompress(base64.b64decode(payload)).decode("utf-8"))
    except (ValueError, zlib.error):
        raise DataFileError("Файл повреждён.")
    return _parse_data_dict(data)


def load_config(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            cfg = json.load(f)
        if isinstance(cfg, dict):
            return cfg
    except (OSError, ValueError):
        pass
    return {}


def save_config(path, cfg):
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=1)
    except OSError:
        pass


# ----------------------------------------------------------------------------
# Экспорт: форма «как в Excel»
# ----------------------------------------------------------------------------
# Повторяет исходную таблицу: №, Ф.И.О., недельная нагрузка,
# «время ежедневной работы» в виде «5,70+20», время прихода, время ухода,
# подпись. Если у кого-то своё число рабочих дней — добавляется столбец «Дней».

def report_title(settings):
    return "Расчет рабочего времени в период каникул с {}г. по {}г.".format(
        settings.date_from.strftime("%d.%m.%Y"), settings.date_to.strftime("%d.%m.%Y"))


def daily_header(settings):
    return ("Время ежедневной работы на каникулах при {}-дневной рабочей неделе "
            "(обеденный перерыв - {} мин.)".format(settings.days, settings.lunch))


def work_text(res, settings):
    """Как в исходной таблице: '5,70+20' (часы в день без обеда + обед в минутах)."""
    return "{}+{}".format(fmt_decimal(res.work_hours), settings.lunch)


def has_own_days(settings, employees):
    return any(e.days and e.days != settings.days and not e.note for e in employees)


def check_before_export(settings, employees):
    if not employees:
        raise InputError("Список работников пуст.")
    bad = [e.name for e in employees if not e.note and calc_employee(e, settings).is_error]
    if bad:
        raise InputError("У этих работников время ухода получается после полуночи, "
                         "исправьте данные:\n" + "\n".join(bad[:15]))


# ----------------------------------------------------------------------------
# Экспорт в Excel
# ----------------------------------------------------------------------------

def export_excel(path, settings, employees):
    """Колонки как в исходном файле: A №, B Ф.И.О., C нагрузка, D (дней, скрыт),
    E «5,70+20», F (минут работы в день, скрыт), G (скрыт), H приход, I уход, J подпись."""
    check_before_export(settings, employees)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_head = Font(name="Calibri", size=11)
    f_data = Font(name="Times New Roman", size=12)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    show_days = has_own_days(settings, employees)

    ws["A1"] = report_title(settings)
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(name="Calibri", size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 21

    hr = 3
    heads = {
        "A": "№", "B": "Ф.И.О. работника",
        "C": "Оплачиваемая недельная педнагрузка ч.",
        "D": "Рабочих дней в неделю",
        "E": daily_header(settings),
        "H": "Время прихода", "I": "Время ухода", "J": "Подпись пед. работника",
    }
    for col in "ABCDEFGHIJ":
        cell = ws["{}{}".format(col, hr)]
        cell.value = heads.get(col)
        cell.font = f_head
        cell.alignment = center
        cell.border = border
    ws.merge_cells("E{0}:G{0}".format(hr))
    ws.row_dimensions[hr].height = 105.75

    for n, e in enumerate(employees, start=1):
        r = hr + n
        ws["A{}".format(r)] = n
        ws["B{}".format(r)] = e.name
        if e.note:
            # текст на всю строку: от нагрузки до времени ухода включительно
            ws["C{}".format(r)] = e.note
            for col in "ABCDEFGHIJ":
                cell = ws["{}{}".format(col, r)]
                cell.font = f_data
                cell.border = border
                cell.alignment = left if col == "B" else center
            ws.merge_cells("C{0}:I{0}".format(r))
            ws.row_dimensions[r].height = 23.1
            continue
        days = e.days_for(settings)
        ws["C{}".format(r)] = float(e.load)
        ws["D{}".format(r)] = days
        # FIXED учитывает язык Excel (в русском — запятая), в отличие от TEXT(…;"0,00")
        ws["E{}".format(r)] = '=FIXED(C{0}/D{0},2,TRUE)&"+{1}"'.format(r, settings.lunch)
        # минут работы в день, округлено вверх; ROUND(…,6) убирает «хвосты» дробей
        ws["F{}".format(r)] = "=ROUNDUP(ROUND(C{0}*60/D{0},6),0)".format(r)
        ws["H{}".format(r)] = datetime.time(e.arrival // 60, e.arrival % 60)
        ws["I{}".format(r)] = "=H{0}+(F{0}+{1})/1440".format(r, settings.lunch)
        ws["C{}".format(r)].number_format = "0.0#"
        ws["H{}".format(r)].number_format = "h:mm"
        ws["I{}".format(r)].number_format = "h:mm"
        for col in "ABCDEFGHIJ":
            cell = ws["{}{}".format(col, r)]
            cell.font = f_data
            cell.border = border
            cell.alignment = left if col == "B" else center
        ws.row_dimensions[r].height = 23.1

    widths = {"A": 6.86, "B": 38.43, "C": 14.71, "D": 10, "E": 19.71, "F": 6,
              "G": 6, "H": 9.14, "I": 9.86, "J": 13.43}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["D"].hidden = not show_days
    ws.column_dimensions["F"].hidden = True
    ws.column_dimensions["G"].hidden = True

    ws.print_title_rows = "{0}:{0}".format(hr)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.71
    ws.page_margins.top = ws.page_margins.bottom = 0.75

    wb.calculation.fullCalcOnLoad = True  # Excel пересчитает формулы при открытии
    wb.save(path)


# ----------------------------------------------------------------------------
# Экспорт в PDF
# ----------------------------------------------------------------------------

_FONT_CANDIDATES = [
    # (обычный, жирный) — Windows
    ("times.ttf", "timesbd.ttf"),
    ("arial.ttf", "arialbd.ttf"),
    ("tahoma.ttf", "tahomabd.ttf"),
    # Linux (для проверки без Windows)
    ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf"),
    ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
]


def _font_dirs():
    dirs = []
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windir:
        dirs.append(os.path.join(windir, "Fonts"))
    dirs.append(os.path.join(app_dir(), "fonts"))
    dirs += ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/TTF", "/Library/Fonts"]
    return dirs


def _register_fonts():
    """Регистрирует шрифт с кириллицей. Возвращает (обычный, жирный)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "RK-Regular" in pdfmetrics.getRegisteredFontNames():
        return "RK-Regular", "RK-Bold"
    for d in _font_dirs():
        for reg, bold in _FONT_CANDIDATES:
            p_reg, p_bold = os.path.join(d, reg), os.path.join(d, bold)
            if os.path.isfile(p_reg):
                pdfmetrics.registerFont(TTFont("RK-Regular", p_reg))
                pdfmetrics.registerFont(TTFont("RK-Bold", p_bold if os.path.isfile(p_bold) else p_reg))
                return "RK-Regular", "RK-Bold"
    raise InputError("Не найден шрифт с русскими буквами (Times New Roman или Arial) "
                     "в папке Windows\\Fonts.")


def export_pdf(path, settings, employees):
    """Та же форма, что в Excel: видимые столбцы исходной таблицы."""
    check_before_export(settings, employees)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    f_reg, f_bold = _register_fonts()
    st_title = ParagraphStyle("t", fontName=f_bold, fontSize=14, leading=17, alignment=TA_CENTER)
    st_head = ParagraphStyle("h", fontName=f_reg, fontSize=9.5, leading=11.5, alignment=TA_CENTER)
    st_cell = ParagraphStyle("c", fontName=f_reg, fontSize=11, leading=13, alignment=TA_CENTER)
    st_name = ParagraphStyle("n", fontName=f_reg, fontSize=11, leading=13, alignment=TA_LEFT)
    show_days = has_own_days(settings, employees)

    # ширины как в исходном Excel (в символах), пересчитываются под ширину листа
    cols = [("№", 5), ("Ф.И.О. работника", 39), ("Оплачиваемая недельная педнагрузка ч.", 17)]
    if show_days:
        cols.append(("Дней в неделю", 10))
    cols += [(daily_header(settings), 19), ("Время прихода", 10.5), ("Время ухода", 10),
             ("Подпись пед. работника", 14)]

    data = [[Paragraph(h, st_head) for h, _ in cols]]
    spans = []
    for n, e in enumerate(employees, start=1):
        if e.note:
            row = [Paragraph(str(n), st_cell), Paragraph(e.name, st_name),
                   Paragraph(e.note, st_cell)] + [""] * (len(cols) - 3)
            spans.append(("SPAN", (2, n), (len(cols) - 2, n)))
            data.append(row)
            continue
        res = calc_employee(e, settings)
        row = [Paragraph(str(n), st_cell), Paragraph(e.name, st_name),
               Paragraph(fmt_load(e.load), st_cell)]
        if show_days:
            row.append(Paragraph(str(res.days), st_cell))
        row += [Paragraph(work_text(res, settings), st_cell),
                Paragraph(fmt_time(e.arrival), st_cell),
                Paragraph(fmt_time(res.departure), st_cell), ""]
        data.append(row)

    page_w = A4[0] - 24 * mm
    total = sum(w for _, w in cols)
    table = Table(data, colWidths=[page_w * w / total for _, w in cols], repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ] + spans))

    title = report_title(settings)

    doc = SimpleDocTemplate(path, pagesize=A4, title=title, author=APP_TITLE,
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm)
    doc.build([Paragraph(title, st_title), Spacer(1, 5 * mm), table])

# ----------------------------------------------------------------------------
# Диалог добавления / изменения работника
# ----------------------------------------------------------------------------

def _days_default_label(settings):
    return "как в параметрах ({})".format(settings.days)


class EmployeeDialog(tk.Toplevel):
    def __init__(self, master, settings, employee=None, title="Работник", notes=None):
        tk.Toplevel.__init__(self, master)
        self.withdraw()
        self.title(title)
        self.resizable(False, False)
        self.transient(master)
        self.settings = settings
        self.result = None

        frm = ttk.Frame(self, padding=12)
        frm.grid(sticky="nsew")

        self.days_values = [_days_default_label(settings)] + DAYS_CHOICES
        self.var_name = tk.StringVar(value=employee.name if employee else "")
        self.var_load = tk.StringVar(
            value=fmt_load(employee.load) if employee and employee.load is not None else "")
        arr = employee.arrival if employee and employee.arrival is not None else settings.default_arrival
        self.var_arr = tk.StringVar(value=fmt_time(arr))
        self.var_days = tk.StringVar(
            value=str(employee.days) if employee and employee.days else self.days_values[0])
        self.var_mode = tk.StringVar(value="note" if employee and employee.note else "calc")
        self.var_note = tk.StringVar(value=employee.note if employee and employee.note else DEFAULT_NOTE)

        ttk.Label(frm, text="Ф.И.О. работника:").grid(row=0, column=0, sticky="w", pady=3)
        e_name = ttk.Entry(frm, textvariable=self.var_name, width=46)
        e_name.grid(row=0, column=1, sticky="we", pady=3)

        ttk.Radiobutton(frm, text="Расчёт по нагрузке", variable=self.var_mode, value="calc",
                        command=self._mode_changed).grid(row=1, column=0, columnspan=2,
                                                         sticky="w", pady=(10, 2))
        self.calc_widgets = []
        ttk.Label(frm, text="    Недельная нагрузка, ч:").grid(row=2, column=0, sticky="w", pady=3)
        w = ttk.Entry(frm, textvariable=self.var_load, width=10)
        w.grid(row=2, column=1, sticky="w", pady=3)
        self.calc_widgets.append(w)
        ttk.Label(frm, text="    Делить на дней в неделю:").grid(row=3, column=0, sticky="w", pady=3)
        w = ttk.Combobox(frm, textvariable=self.var_days, values=self.days_values, state="readonly", width=22)
        w.grid(row=3, column=1, sticky="w", pady=3)
        self.calc_widgets.append(w)
        ttk.Label(frm, text="    Время прихода (ч:мм):").grid(row=4, column=0, sticky="w", pady=3)
        w = ttk.Entry(frm, textvariable=self.var_arr, width=10)
        w.grid(row=4, column=1, sticky="w", pady=3)
        self.calc_widgets.append(w)

        ttk.Radiobutton(frm, text="Текст на всю строку вместо расчёта:", variable=self.var_mode,
                        value="note", command=self._mode_changed).grid(row=5, column=0, columnspan=2,
                                                                       sticky="w", pady=(10, 2))
        note_values = [DEFAULT_NOTE] + sorted(set(n for n in (notes or []) if n and n != DEFAULT_NOTE))
        self.cb_note = ttk.Combobox(frm, textvariable=self.var_note, values=note_values, width=52)
        self.cb_note.grid(row=6, column=0, columnspan=2, sticky="we", pady=3, padx=(18, 0))

        self.lbl_info = ttk.Label(frm, text="", foreground="#1a5e1a", justify="left")
        self.lbl_info.grid(row=7, column=0, columnspan=2, sticky="w", pady=(10, 2))
        self.lbl_err = ttk.Label(frm, text="", foreground="#b00020", justify="left", wraplength=420)
        self.lbl_err.grid(row=8, column=0, columnspan=2, sticky="w")

        btns = ttk.Frame(frm)
        btns.grid(row=9, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(btns, text="Сохранить", command=self.on_ok).pack(side="left", padx=4)
        ttk.Button(btns, text="Отмена", command=self.destroy).pack(side="left")

        for v in (self.var_name, self.var_load, self.var_arr, self.var_days, self.var_note):
            v.trace_add("write", lambda *a: self.update_preview())
        self.bind("<Return>", lambda e: self.on_ok())
        self.bind("<Escape>", lambda e: self.destroy())

        self._mode_changed()
        self.update_idletasks()
        x = master.winfo_rootx() + (master.winfo_width() - self.winfo_reqwidth()) // 2
        y = master.winfo_rooty() + (master.winfo_height() - self.winfo_reqheight()) // 3
        self.geometry("+{}+{}".format(max(x, 0), max(y, 0)))
        self.deiconify()
        self.grab_set()
        e_name.focus_set()
        e_name.icursor("end")

    def _is_note(self):
        return self.var_mode.get() == "note"

    def _mode_changed(self):
        note = self._is_note()
        for w in self.calc_widgets:
            if isinstance(w, ttk.Combobox):
                w.configure(state="disabled" if note else "readonly")
            else:
                w.configure(state="disabled" if note else "normal")
        self.cb_note.configure(state="normal" if note else "disabled")
        self.update_preview()

    def _days(self):
        v = self.var_days.get()
        return None if v == self.days_values[0] else parse_days(v)

    def _calc(self):
        load = parse_load(self.var_load.get())
        arr = parse_time(self.var_arr.get())
        days = self._days()
        res = calculate(load, arr, self.settings, days)
        if res.is_error:
            raise InputError("Время ухода получается {} — позже полуночи. "
                             "Проверьте нагрузку и время прихода.".format(fmt_time(res.departure)))
        return load, arr, days, res

    def update_preview(self):
        if self._is_note():
            try:
                note = parse_note(self.var_note.get())
                self.lbl_info.config(text="В таблице вместо нагрузки и времени будет:\n«{}»".format(note))
                self.lbl_err.config(text="")
            except InputError as e:
                self.lbl_info.config(text="")
                self.lbl_err.config(text=str(e))
            return
        try:
            _, _, _, res = self._calc()
            self.lbl_info.config(text="В день: {} / {} = {} ч → {} (ч:мм) + обед {} мин\nУход: {}".format(
                fmt_load(Decimal(self.var_load.get().replace(",", "."))), res.days,
                fmt_decimal(res.work_hours), fmt_duration(res.work_min),
                self.settings.lunch, fmt_time(res.departure)))
            self.lbl_err.config(text="")
            if res.is_late:
                self.lbl_err.config(text="Внимание: уход позже {}.".format(fmt_time(LATE_WARNING_MIN)))
        except (InputError, InvalidOperation) as e:
            self.lbl_info.config(text="")
            self.lbl_err.config(text=str(e))

    def on_ok(self):
        try:
            name = parse_name(self.var_name.get())
            if self._is_note():
                note = parse_note(self.var_note.get())
                # нагрузку и приход сохраняем, если они введены верно, — пригодятся,
                # если потом вернуть расчёт
                try:
                    load = parse_load(self.var_load.get())
                except InputError:
                    load = None
                try:
                    arr = parse_time(self.var_arr.get())
                except InputError:
                    arr = None
                self.result = Employee(name, load, arr, self._days(), note)
            else:
                load, arr, days, _ = self._calc()
                self.result = Employee(name, load, arr, days)
        except InputError as e:
            self.lbl_err.config(text=str(e))
            self.bell()
            return
        self.destroy()


# ----------------------------------------------------------------------------
# Главное окно
# ----------------------------------------------------------------------------

COLUMNS = [
    ("n", "№", 45, "center"),
    ("name", "Ф.И.О. работника", 300, "w"),
    ("load", "Нагрузка, ч/нед", 140, "center"),
    ("days", "Дней", 70, "center"),
    ("work", "Время ежедн. работы", 190, "center"),
    ("arr", "Приход", 90, "center"),
    ("dep", "Уход", 90, "center"),
]


class App(tk.Tk):
    def __init__(self):
        tk.Tk.__init__(self)
        self.title(APP_TITLE)
        self.minsize(950, 520)
        self.geometry("1060x660")

        self.default_path = os.path.join(data_dir(), DATA_FILE_NAME)
        self.config_path = os.path.join(data_dir(), CONFIG_FILE_NAME)
        self.app_cfg = load_config(self.config_path)
        self.data_path = self.default_path      # текущий файл, в него идёт автосохранение
        self.settings = Settings()
        self.employees = []
        self._loading = True
        self._save_blocked = False   # не затирать файл, который не удалось прочитать

        self._build_ui()
        self._load_initial()
        self._loading = False
        self.refresh()
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------- интерфейс ----------

    def _build_ui(self):
        style = ttk.Style(self)
        try:
            if "vista" in style.theme_names():
                style.theme_use("vista")
        except tk.TclError:
            pass
        style.configure("Treeview", rowheight=22)

        # Меню
        menubar = tk.Menu(self)
        m_file = tk.Menu(menubar, tearoff=0)
        m_file.add_command(label="Открыть список…", command=self.on_open_list)
        self.m_recent = tk.Menu(m_file, tearoff=0)
        m_file.add_cascade(label="Недавние списки", menu=self.m_recent)
        m_file.add_command(label="Сохранить список как…", command=self.on_save_list)
        m_file.add_separator()
        m_file.add_command(label="Экспорт в Excel…", command=self.on_export_excel)
        m_file.add_command(label="Экспорт в PDF…", command=self.on_export_pdf)
        m_file.add_separator()
        m_file.add_command(label="Выход", command=self.on_close)
        menubar.add_cascade(label="Файл", menu=m_file)
        self.config(menu=menubar)

        # Параметры
        pf = ttk.LabelFrame(self, text="Параметры расчёта", padding=8)
        pf.pack(fill="x", padx=8, pady=(8, 4))

        self.var_from = tk.StringVar()
        self.var_to = tk.StringVar()
        self.var_days = tk.StringVar()
        self.var_lunch = tk.StringVar()
        self.var_default_arr = tk.StringVar()

        ttk.Label(pf, text="Каникулы с:").grid(row=0, column=0, sticky="w")
        ttk.Entry(pf, textvariable=self.var_from, width=12).grid(row=0, column=1, padx=(4, 12), sticky="w")
        ttk.Label(pf, text="по:").grid(row=0, column=2, sticky="w")
        ttk.Entry(pf, textvariable=self.var_to, width=12).grid(row=0, column=3, padx=(4, 12), sticky="w")
        ttk.Label(pf, text="Обед, мин:").grid(row=0, column=4, sticky="e")
        ttk.Spinbox(pf, from_=0, to=180, increment=5, textvariable=self.var_lunch, width=5).grid(
            row=0, column=5, padx=(4, 12), sticky="w")
        ttk.Label(pf, text="Приход по умолчанию:").grid(row=0, column=6, sticky="e")
        ttk.Entry(pf, textvariable=self.var_default_arr, width=6).grid(row=0, column=7, padx=(4, 0), sticky="w")

        ttk.Label(pf, text="Делить недельную нагрузку на:").grid(row=1, column=0, columnspan=2,
                                                                 sticky="w", pady=(8, 0))
        ttk.Combobox(pf, textvariable=self.var_days, values=DAYS_CHOICES, state="readonly",
                     width=4).grid(row=1, column=2, sticky="w", pady=(8, 0))
        self.lbl_days_word = ttk.Label(pf, text="дней в неделю (для всех, у кого не указано своё)")
        self.lbl_days_word.grid(row=1, column=3, columnspan=5, sticky="w", pady=(8, 0), padx=(4, 0))

        ttk.Label(pf, text="Время работы в день округляется вверх до целой минуты.",
                  foreground="#555555").grid(row=2, column=0, columnspan=8, sticky="w", pady=(6, 0))
        self.lbl_settings_err = ttk.Label(pf, text="", foreground="#b00020")
        self.lbl_settings_err.grid(row=3, column=0, columnspan=8, sticky="w", pady=(2, 0))

        for v in (self.var_from, self.var_to, self.var_days, self.var_lunch, self.var_default_arr):
            v.trace_add("write", lambda *a: self.on_settings_changed())

        # Кнопки
        tb = ttk.Frame(self, padding=(8, 2))
        tb.pack(fill="x")
        for text, cmd in [("Добавить", self.on_add), ("Изменить", self.on_edit),
                          ("Удалить", self.on_delete),
                          ("Приход / дни для выделенных…", self.on_bulk_edit),
                          ("Сортировать по Ф.И.О.", self.on_sort)]:
            ttk.Button(tb, text=text, command=cmd).pack(side="left", padx=(0, 4))
        ttk.Button(tb, text="Экспорт в PDF…", command=self.on_export_pdf).pack(side="right")
        ttk.Button(tb, text="Экспорт в Excel…", command=self.on_export_excel).pack(side="right", padx=4)

        # Таблица
        tf = ttk.Frame(self)
        tf.pack(fill="both", expand=True, padx=8, pady=4)
        self.tree = ttk.Treeview(tf, columns=[c[0] for c in COLUMNS], show="headings",
                                 selectmode="extended")
        for cid, title, width, anchor in COLUMNS:
            self.tree.heading(cid, text=title)
            self.tree.column(cid, width=width, anchor=anchor, stretch=(cid == "name"))
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self.tree.tag_configure("late", background="#fff3c4")
        self.tree.tag_configure("error", background="#ffd0d0")
        self.tree.tag_configure("note", foreground="#1f4e9a")

        self.tree.bind("<Double-1>", lambda e: self.on_edit())
        self.tree.bind("<Return>", lambda e: self.on_edit())
        self.tree.bind("<Delete>", lambda e: self.on_delete())
        self.bind("<Insert>", lambda e: self.on_add())
        self.bind("<Control-n>", lambda e: self.on_add())
        self.bind("<Control-a>", lambda e: self.tree.selection_set(self.tree.get_children()))

        self.status = ttk.Label(self, text="", anchor="w", padding=(8, 2))
        self.status.pack(fill="x", side="bottom")

    # ---------- данные ----------

    # ---------- текущий файл / недавние ----------

    def _set_current(self, path):
        """Делает файл текущим: в него автосохранение, его откроем при следующем запуске."""
        path = os.path.abspath(path)
        self.data_path = path
        recent = [p for p in self.app_cfg.get("recent", [])
                  if isinstance(p, str) and os.path.normcase(p) != os.path.normcase(path)]
        self.app_cfg["last_file"] = path
        self.app_cfg["recent"] = [path] + recent[:MAX_RECENT - 1]
        save_config(self.config_path, self.app_cfg)
        self._update_title()
        self._build_recent_menu()

    def _update_title(self):
        name = os.path.basename(self.data_path)
        self.title("{} — {}".format(name, APP_TITLE))

    def _build_recent_menu(self):
        self.m_recent.delete(0, "end")
        recent = [p for p in self.app_cfg.get("recent", []) if isinstance(p, str)]
        if not recent:
            self.m_recent.add_command(label="(пусто)", state="disabled")
            return
        for i, p in enumerate(recent, start=1):
            label = "{}  {}".format(i, p)
            if not os.path.exists(p):
                label += "  (не найден)"
            self.m_recent.add_command(label=label, command=lambda p=p: self.open_list(p))

    def _load_initial(self):
        last = self.app_cfg.get("last_file")
        if isinstance(last, str) and last:
            if os.path.exists(last):
                self.data_path = last
            else:
                messagebox.showwarning(
                    APP_TITLE, "Последний открытый список не найден:\n{}\n\n"
                               "Будет открыт основной список.".format(last))
        self._load_file_at_start()
        self._set_current(self.data_path)

    def _load_file_at_start(self):
        bad = 0
        if os.path.exists(self.data_path):
            try:
                self.settings, self.employees, bad = load_data(self.data_path)
            except DataFileError as e:
                self._save_blocked = True
                messagebox.showerror(
                    APP_TITLE, "{}\n\nФайл: {}\n\nСписок не загружен. Чтобы не потерять файл, "
                               "автосохранение отключено до тех пор, пока вы не откроете "
                               "другой список (Файл → Открыть)."
                               .format(e, self.data_path))
        self._settings_to_vars()
        if bad:
            messagebox.showwarning(APP_TITLE, "{} записей содержали ошибки и пропущены.".format(bad))

    def _settings_to_vars(self):
        s = self.settings
        self.var_from.set(s.date_from.strftime("%d.%m.%Y"))
        self.var_to.set(s.date_to.strftime("%d.%m.%Y"))
        self.var_days.set(str(s.days))
        self.var_lunch.set(str(s.lunch))
        self.var_default_arr.set(fmt_time(s.default_arrival))

    def _read_settings(self):
        d_from = parse_date(self.var_from.get(), "с")
        d_to = parse_date(self.var_to.get(), "по")
        if d_to < d_from:
            raise InputError("Дата окончания каникул раньше даты начала.")
        days = parse_days(self.var_days.get())
        lunch = parse_int_range(self.var_lunch.get(), "Обед", 0, 180)
        try:
            arr = parse_time(self.var_default_arr.get())
        except InputError as e:
            raise InputError("Приход по умолчанию: {}".format(e))
        return Settings(d_from, d_to, days, lunch, arr)

    def settings_valid(self):
        try:
            self._read_settings()
            return True
        except InputError as e:
            self.lbl_settings_err.config(text="Ошибка в параметрах: {}".format(e))
            messagebox.showerror(APP_TITLE, "Исправьте ошибку в параметрах расчёта:\n{}".format(e))
            return False

    def on_settings_changed(self):
        if self._loading:
            return
        try:
            self.settings = self._read_settings()
            self.lbl_settings_err.config(text="")
        except InputError as e:
            self.lbl_settings_err.config(
                text="Ошибка в параметрах: {} (расчёт по последним верным значениям)".format(e))
            return
        self.lbl_days_word.config(text="{} в неделю (для всех, у кого не указано своё)".format(
            days_word(self.settings.days)))
        self.refresh()
        self.autosave()

    def autosave(self):
        if self._save_blocked:
            return
        try:
            save_data(self.data_path, self.settings, self.employees)
        except OSError as e:
            self.status.config(text="Не удалось сохранить данные: {}".format(e))

    def refresh(self, select=None):
        selected = set(select if select is not None else [int(i) for i in self.tree.selection()])
        self.tree.delete(*self.tree.get_children())
        n_late = n_err = 0
        n_note = 0
        for i, e in enumerate(self.employees):
            res = calc_employee(e, self.settings)
            if res is None:
                n_note += 1
                self.tree.insert("", "end", iid=str(i), tags=("note",), values=(
                    i + 1, e.name, "", "", e.note, "", ""))
                continue
            tags = ()
            if res.is_error:
                tags = ("error",)
                n_err += 1
                dep = "ОШИБКА"
            else:
                dep = fmt_time(res.departure)
                if res.is_late:
                    tags = ("late",)
                    n_late += 1
            days_txt = str(res.days) + (" *" if e.days else "")
            self.tree.insert("", "end", iid=str(i), tags=tags, values=(
                i + 1, e.name, fmt_load(e.load), days_txt, work_text(res, self.settings),
                fmt_time(e.arrival), dep))
        sel = [str(i) for i in selected if 0 <= i < len(self.employees)]
        if sel:
            self.tree.selection_set(sel)
            self.tree.see(sel[0])
        msg = "Работников: {}".format(len(self.employees))
        if n_note:
            msg += "   |   с текстом вместо расчёта: {}".format(n_note)
        if any(e.days and not e.note for e in self.employees):
            msg += "   |   * — своё число дней"
        if n_late:
            msg += "   |   уход позже {}: {} (жёлтые)".format(fmt_time(LATE_WARNING_MIN), n_late)
        if n_err:
            msg += "   |   ОШИБКИ (уход после полуночи): {} (красные)".format(n_err)
        if self._save_blocked:
            msg += "   |   АВТОСОХРАНЕНИЕ ОТКЛЮЧЕНО"
        else:
            msg += "   |   Файл: {}".format(self.data_path)
        self.status.config(text=msg)

    def _notes(self):
        return [e.note for e in self.employees if e.note]

    def _selected_indices(self):
        return sorted(int(i) for i in self.tree.selection())

    def _name_exists(self, name, skip_index=None):
        key = name.lower().replace("ё", "е")
        return any(i != skip_index and e.name.lower().replace("ё", "е") == key
                   for i, e in enumerate(self.employees))

    def _replace_all(self, settings, employees):
        self._loading = True
        if settings is not None:
            self.settings = settings
            self._settings_to_vars()
        self._loading = False
        self.lbl_settings_err.config(text="")
        self.employees = employees
        self._save_blocked = False
        self.refresh(select=[])
        self.autosave()

    # ---------- действия ----------

    def on_add(self):
        if not self.settings_valid():
            return
        dlg = EmployeeDialog(self, self.settings, title="Добавить работника", notes=self._notes())
        self.wait_window(dlg)
        if dlg.result is None:
            return
        if self._name_exists(dlg.result.name) and not messagebox.askyesno(
                APP_TITLE, "Работник «{}» уже есть в списке. Всё равно добавить?".format(dlg.result.name)):
            return
        self.employees.append(dlg.result)
        self.refresh(select=[len(self.employees) - 1])
        self.autosave()

    def on_edit(self):
        idx = self._selected_indices()
        if not idx:
            messagebox.showinfo(APP_TITLE, "Выберите работника в таблице.")
            return
        if not self.settings_valid():
            return
        i = idx[0]
        dlg = EmployeeDialog(self, self.settings, self.employees[i], title="Изменить данные работника",
                             notes=self._notes())
        self.wait_window(dlg)
        if dlg.result is None:
            return
        if self._name_exists(dlg.result.name, skip_index=i) and not messagebox.askyesno(
                APP_TITLE, "Работник «{}» уже есть в списке. Всё равно сохранить?".format(dlg.result.name)):
            return
        self.employees[i] = dlg.result
        self.refresh(select=[i])
        self.autosave()

    def on_delete(self):
        idx = self._selected_indices()
        if not idx:
            messagebox.showinfo(APP_TITLE, "Выберите работников для удаления.")
            return
        q = ("Удалить «{}»?".format(self.employees[idx[0]].name) if len(idx) == 1
             else "Удалить выбранных работников ({})?".format(len(idx)))
        if not messagebox.askyesno(APP_TITLE, q):
            return
        for i in reversed(idx):
            del self.employees[i]
        self.refresh(select=[])
        self.autosave()

    def on_bulk_edit(self):
        idx = self._selected_indices()
        if not idx:
            messagebox.showinfo(APP_TITLE, "Выделите работников (Ctrl или Shift + щелчок, Ctrl+A — всех).")
            return
        if not self.settings_valid():
            return
        win = tk.Toplevel(self)
        win.title("Изменить для выделенных ({})".format(len(idx)))
        win.resizable(False, False)
        win.transient(self)
        frm = ttk.Frame(win, padding=12)
        frm.pack()

        keep = "не менять"
        var_chk_arr = tk.BooleanVar(value=True)
        var_arr = tk.StringVar(value=fmt_time(self.settings.default_arrival))
        days_values = [keep, _days_default_label(self.settings)] + DAYS_CHOICES
        var_days = tk.StringVar(value=keep)

        ttk.Checkbutton(frm, text="Время прихода (ч:мм):", variable=var_chk_arr).grid(row=0, column=0, sticky="w")
        ent = ttk.Entry(frm, textvariable=var_arr, width=8)
        ent.grid(row=0, column=1, padx=6, sticky="w")
        ttk.Label(frm, text="Делить на дней в неделю:").grid(row=1, column=0, sticky="w", pady=(6, 0))
        ttk.Combobox(frm, textvariable=var_days, values=days_values, state="readonly", width=22).grid(
            row=1, column=1, padx=6, sticky="w", pady=(6, 0))
        err = ttk.Label(frm, text="", foreground="#b00020", wraplength=380)
        err.grid(row=2, column=0, columnspan=2, sticky="w", pady=4)

        def apply(_e=None):
            new_arr = None
            if var_chk_arr.get():
                try:
                    new_arr = parse_time(var_arr.get())
                except InputError as e:
                    err.config(text=str(e))
                    return
            dv = var_days.get()
            change_days = dv != keep
            new_days = None if dv == days_values[1] else (parse_days(dv) if change_days else None)
            bad = []
            for i in idx:
                e = self.employees[i]
                if e.note:
                    continue
                arr = e.arrival if new_arr is None else new_arr
                days = new_days if change_days else e.days
                if calculate(e.load, arr, self.settings, days).is_error:
                    bad.append(e.name)
            if bad:
                err.config(text="Уход получится после полуночи у: " + ", ".join(bad[:3]))
                return
            for i in idx:
                if new_arr is not None:
                    self.employees[i].arrival = new_arr
                if change_days:
                    self.employees[i].days = new_days
            win.destroy()
            self.refresh(select=idx)
            self.autosave()

        b = ttk.Frame(frm)
        b.grid(row=3, column=0, columnspan=2, sticky="e")
        ttk.Button(b, text="Применить", command=apply).pack(side="left", padx=4)
        ttk.Button(b, text="Отмена", command=win.destroy).pack(side="left")
        win.bind("<Return>", apply)
        win.bind("<Escape>", lambda e: win.destroy())
        win.update_idletasks()
        win.geometry("+{}+{}".format(self.winfo_rootx() + 150, self.winfo_rooty() + 150))
        win.grab_set()
        ent.focus_set()
        ent.select_range(0, "end")

    def on_sort(self):
        self.employees.sort(key=lambda e: e.name.lower().replace("ё", "е"))
        self.refresh(select=[])
        self.autosave()

    def on_open_list(self):
        path = filedialog.askopenfilename(
            title="Открыть список", initialdir=os.path.dirname(self.data_path) or data_dir(),
            filetypes=[("Список работников", "*" + DATA_EXT), ("Все файлы", "*.*")])
        if path:
            self.open_list(path)

    def open_list(self, path):
        """Открывает список. Текущий уже сохранён автоматически, поэтому ничего не теряется."""
        if not os.path.exists(path):
            messagebox.showerror(APP_TITLE, "Файл не найден:\n{}".format(path))
            self._build_recent_menu()
            return
        try:
            settings, emps, bad = load_data(path)
        except DataFileError as e:
            messagebox.showerror(APP_TITLE, "{}\n\nФайл: {}".format(e, path))
            return
        self.autosave()
        self._set_current(path)
        self._replace_all(settings, emps)
        if bad:
            messagebox.showwarning(APP_TITLE, "{} записей содержали ошибки и пропущены.".format(bad))

    def on_save_list(self):
        default = "Список_{}{}".format(self.settings.date_from.strftime("%d.%m.%Y"), DATA_EXT)
        path = filedialog.asksaveasfilename(
            title="Сохранить список", defaultextension=DATA_EXT, initialfile=default,
            filetypes=[("Список работников", "*" + DATA_EXT)])
        if not path:
            return
        try:
            save_data(path, self.settings, self.employees)
        except OSError as e:
            messagebox.showerror(APP_TITLE, "Не удалось сохранить:\n{}".format(e))
            return
        self._save_blocked = False
        self._set_current(path)
        self.refresh()
        messagebox.showinfo(APP_TITLE, "Список сохранён:\n{}\n\nДальше изменения сохраняются "
                                       "в этот файл, и он откроется при следующем запуске.".format(path))

    def _export(self, kind):
        if not self.settings_valid():
            return
        try:
            check_before_export(self.settings, self.employees)
        except InputError as e:
            messagebox.showerror(APP_TITLE, str(e))
            return
        ext, label, func = {
            "xlsx": (".xlsx", "Книга Excel", export_excel),
            "pdf": (".pdf", "Документ PDF", export_pdf),
        }[kind]
        default = "Расчет_рабочего_времени_{}-{}{}".format(
            self.settings.date_from.strftime("%d.%m.%Y"), self.settings.date_to.strftime("%d.%m.%Y"), ext)
        path = filedialog.asksaveasfilename(
            title="Сохранить расчёт", defaultextension=ext, initialfile=default,
            filetypes=[(label, "*" + ext)])
        if not path:
            return
        try:
            func(path, self.settings, self.employees)
        except PermissionError:
            messagebox.showerror(APP_TITLE, "Не удалось записать файл. Возможно, он открыт в другой "
                                            "программе — закройте его и повторите.")
            return
        except InputError as e:
            messagebox.showerror(APP_TITLE, str(e))
            return
        except Exception as e:
            messagebox.showerror(APP_TITLE, "Ошибка при сохранении:\n{}".format(e))
            return
        if messagebox.askyesno(APP_TITLE, "Файл сохранён:\n{}\n\nОткрыть его?".format(path)):
            try:
                os.startfile(path)  # только Windows
            except (AttributeError, OSError) as e:
                messagebox.showwarning(APP_TITLE, "Не удалось открыть файл: {}".format(e))

    def on_export_excel(self):
        self._export("xlsx")

    def on_export_pdf(self):
        self._export("pdf")

    def on_close(self):
        self.autosave()
        self.destroy()


def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
